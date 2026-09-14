"""
VideoMind AI: Streamlit web application.

Run with:
streamlit run app.py
"""

import os
import io

import streamlit as st
from dotenv import load_dotenv

from rag_pipeline import RAGPipeline
from transcript_loader import parse_uploaded_transcript, register_uploaded_transcript

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document


load_dotenv()


# --------------------------------------------------
# PAGE CONFIGURATION
# --------------------------------------------------

st.set_page_config(
    page_title="VideoMind AI",
    layout="wide",
    initial_sidebar_state="expanded"
)


# --------------------------------------------------
# CUSTOM CSS
# --------------------------------------------------

st.markdown(
    """
    <style>
        .main {
            padding: 2rem;
        }

        .stTabs [data-baseweb="tab-list"] button {
            font-size: 1.1rem;
            padding: 0.5rem 1rem;
        }

        .success-box {
            background-color: #d4edda;
            border: 1px solid #c3e6cb;
            border-radius: 5px;
            padding: 1rem;
            color: #155724;
        }

        .error-box {
            background-color: #f8d7da;
            border: 1px solid #f5c6cb;
            border-radius: 5px;
            padding: 1rem;
            color: #721c24;
        }
    </style>
    """,
    unsafe_allow_html=True
)


# --------------------------------------------------
# INITIALIZE SESSION STATE
# --------------------------------------------------

if "rag_pipeline" not in st.session_state:
    try:
        st.session_state.rag_pipeline = RAGPipeline()
    except Exception as error:
        st.error(f"Failed to initialize VideoMind AI: {error}")
        st.stop()


if "processed_videos" not in st.session_state:
    st.session_state.processed_videos = {}


pipeline = st.session_state.rag_pipeline


# --------------------------------------------------
# HELPER FUNCTIONS
# --------------------------------------------------

def _clean_inline_markdown(text: str) -> str:
    """Convert common Markdown formatting into ReportLab-compatible markup."""
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    text = __import__("re").sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = __import__("re").sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = __import__("re").sub(r"`(.+?)`", r"<font name='Courier'>\1</font>", text)
    return text


def create_pdf(notes: str) -> bytes:
    """Create a PDF that keeps the same headings, bullets, and emphasis as the app."""
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50,
    )
    styles = getSampleStyleSheet()
    styles["Title"].fontSize = 20
    styles["Title"].leading = 24
    styles["Heading1"].fontSize = 16
    styles["Heading1"].leading = 20
    styles["Heading2"].fontSize = 13
    styles["Heading2"].leading = 17
    styles["Heading3"].fontSize = 11
    styles["Heading3"].leading = 14
    styles["BodyText"].leading = 15

    story = [Paragraph("VideoMind AI Study Notes", styles["Title"]), Spacer(1, 12)]

    for raw_line in notes.splitlines():
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 6))
            continue

        if line.startswith("### "):
            story.append(Paragraph(_clean_inline_markdown(line[4:]), styles["Heading3"]))
        elif line.startswith("## "):
            story.append(Paragraph(_clean_inline_markdown(line[3:]), styles["Heading2"]))
        elif line.startswith("# "):
            story.append(Paragraph(_clean_inline_markdown(line[2:]), styles["Heading1"]))
        elif line.startswith(("- ", "* ")):
            story.append(Paragraph("• " + _clean_inline_markdown(line[2:]), styles["BodyText"]))
        elif line.startswith("> "):
            story.append(Paragraph(_clean_inline_markdown(line[2:]), styles["BodyText"]))
        else:
            story.append(Paragraph(_clean_inline_markdown(line), styles["BodyText"]))

    document.build(story)
    return buffer.getvalue()


def create_docx(notes: str) -> bytes:
    """Create a Word document with headings, bullets, and bold/italic Markdown."""
    import re

    buffer = io.BytesIO()
    document = Document()
    document.add_heading("VideoMind AI Study Notes", level=1)

    for raw_line in notes.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
            continue
        if line.startswith("## "):
            document.add_heading(line[3:], level=2)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
            continue
        if line.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            content = line[2:]
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            content = line[2:]
        else:
            paragraph = document.add_paragraph()
            content = line

        # Preserve the most common Markdown emphasis used by the app.
        parts = re.split(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", content)
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run(part)
            if part.startswith("**") and part.endswith("**"):
                run.text = part[2:-2]
                run.bold = True
            elif part.startswith("*") and part.endswith("*"):
                run.text = part[1:-1]
                run.italic = True
            elif part.startswith("`") and part.endswith("`"):
                run.text = part[1:-1]
                run.font.name = "Courier New"

    document.save(buffer)
    return buffer.getvalue()

def get_processed_videos():
    """Return all videos stored in the database."""

    try:
        return pipeline.database.list_videos()
    except Exception as error:
        st.error(f"Could not load processed videos: {error}")
        return []


def get_video_title(video):
    """Safely get a display title for a video."""

    return (
        video.get("title")
        or video.get("video_title")
        or video.get("video_id")
        or "Untitled Video"
    )


def get_video_id(video):
    """Safely get the video ID."""

    return video.get("video_id", "")


def show_result_error(result):
    """Display an error returned by the pipeline."""

    error_message = result.get(
        "error",
        "An unknown error occurred."
    )

    st.error(f"Error: {error_message}")


# --------------------------------------------------
# HEADER
# --------------------------------------------------

st.title("VideoMind AI")
st.caption("Learn from YouTube videos with AI.")


st.divider()


# --------------------------------------------------
# SIDEBAR
# --------------------------------------------------

with st.sidebar:
    st.header("About")
    st.caption("Transcript search, AI answers, summaries, and study notes.")

    st.divider()

    st.warning("Use videos up to 15 minutes. About 3–5 videos per day is recommended.")

    with st.expander("Usage details"):
        st.caption("These are recommendations, not guaranteed quotas. Long videos and repeated AI requests may hit Groq limits.")

    st.subheader("Processed Videos")
    videos = get_processed_videos()

    if videos:
        for video in videos:
            video_id = get_video_id(video)
            video_title = get_video_title(video)

            with st.expander(video_title):
                st.caption(video_id)
                if "chunk_count" in video:
                    st.caption(f"Chunks: {video['chunk_count']}")
                if "processed_at" in video:
                    st.caption(f"Processed: {video['processed_at']}")
    else:
        st.caption("No videos processed yet.")


# --------------------------------------------------
# MAIN TABS
# --------------------------------------------------

tab1, tab2, tab3, tab4 = st.tabs(
    [
        "Process Video",
        "Ask Questions",
        "Summary",
        "Study Notes"
    ]
)


# ==================================================
# TAB 1: PROCESS VIDEO
# ==================================================

with tab1:

    st.header("Process a YouTube Video")

    st.caption("Paste a YouTube URL to extract its transcript and prepare it for questions, summaries, and notes.")
    st.info("Recommended: videos up to 15 minutes. About 3–5 videos per day.")

    video_url = st.text_input(
        "YouTube URL",
        placeholder="https://www.youtube.com/watch?v=...",
        key="video_url_input"
    )

    uploaded_file = st.file_uploader(
        "Or upload a transcript (.txt or .vtt)",
        type=["txt", "vtt"],
        key="transcript_file"
    )

    st.info(
        "Good test videos include educational videos, TED Talks, "
        "and videos with English captions."
    )

    if st.button(
        "Process Video",
        use_container_width=True,
        key="process_btn"
    ):

        if not video_url.strip() and uploaded_file is None:
            st.error("Enter a YouTube URL or upload a transcript file.")

        else:

            progress_bar = st.progress(0)
            status_text = st.empty()

            try:
                status_text.text("Loading transcript...")
                progress_bar.progress(20)

                if uploaded_file is not None:
                    transcript_text = parse_uploaded_transcript(
                        uploaded_file.name, uploaded_file.getvalue()
                    )
                    if not transcript_text:
                        raise ValueError("The uploaded transcript is empty.")
                    register_uploaded_transcript("uploaded123", transcript_text)
                    result = pipeline.process_video("uploaded123")
                else:
                    result = pipeline.process_video(video_url)

                progress_bar.progress(100)

                if result.get("success"):

                    st.success("Video processed successfully.")

                    col1, col2, col3 = st.columns(3)

                    with col1:
                        st.metric(
                            "Video ID",
                            result["video_id"][:12] + "..."
                        )

                    with col2:
                        st.metric(
                            "Chunks Created",
                            result["chunk_count"]
                        )

                    with col3:
                        st.metric(
                            "Embeddings",
                            "Generated"
                        )

                    st.session_state.processed_videos[
                        result["video_id"]
                    ] = video_url

                    st.info(
                        "You can now ask questions or generate "
                        "a summary and study notes."
                    )

                else:
                    show_result_error(result)

            except Exception as error:
                st.error(f"Unexpected error: {error}")

            finally:
                progress_bar.empty()
                status_text.empty()


# ==================================================
# TAB 2: ASK QUESTIONS
# ==================================================

with tab2:

    st.header("Ask a Question")

    st.caption("Ask questions about a processed video.")

    videos = get_processed_videos()

    video_options = {
        get_video_id(video): get_video_title(video)
        for video in videos
        if get_video_id(video)
    }

    if not video_options:

        st.warning(
            "No videos processed yet. Go to the Process Video tab first."
        )

    else:

        selected_video_id = st.selectbox(
            "Select a video",
            options=list(video_options.keys()),
            format_func=lambda video_id: video_options[video_id],
            key="video_select"
        )

        st.info(
            f"Selected video: {video_options[selected_video_id]}"
        )

        question = st.text_area(
            "Your Question",
            placeholder=(
                "Example: What are the key takeaways from this video?"
            ),
            height=100,
            key="question_input"
        )

        if st.button(
            "Get Answer",
            use_container_width=True,
            key="answer_btn"
        ):

            if not question.strip():

                st.error("Please enter a question.")

            else:

                try:
                    max_questions = int(
                        os.getenv("MAX_QUESTIONS_PER_DAY", 20)
                    )

                    questions_today = 0

                    if hasattr(
                        pipeline.database,
                        "count_questions_today"
                    ):
                        questions_today = (
                            pipeline.database.count_questions_today()
                        )

                    if questions_today >= max_questions:

                        st.error(
                            f"Daily limit reached "
                            f"({max_questions} questions per day)."
                        )

                    else:

                        with st.spinner(
                            "Searching transcript and generating answer..."
                        ):

                            result = pipeline.answer_question(
                                selected_video_id,
                                question
                            )

                        if result.get("success"):

                            st.success("Answer generated.")

                            st.markdown(result["answer"])

                            st.divider()

                            col1, col2 = st.columns(2)

                            with col1:
                                st.caption(
                                    f"Source chunks: "
                                    f"{result.get('retrieved_chunks', 0)}"
                                )

                            with col2:
                                st.caption(
                                    f"Questions today: "
                                    f"{questions_today + 1}/{max_questions}"
                                )

                        else:
                            show_result_error(result)

                except Exception as error:
                    st.error(f"Unexpected error: {error}")


# ==================================================
# TAB 3: SUMMARY
# ==================================================

with tab3:

    st.header("Generate Summary")

    st.caption("Generate a concise summary of a processed video.")

    videos = get_processed_videos()

    video_options = {
        get_video_id(video): get_video_title(video)
        for video in videos
        if get_video_id(video)
    }

    if not video_options:

        st.warning(
            "No videos processed yet. Go to the Process Video tab first."
        )

    else:

        selected_video_id = st.selectbox(
            "Select a video",
            options=list(video_options.keys()),
            format_func=lambda video_id: video_options[video_id],
            key="summary_video_select"
        )

        if st.button(
            "Generate Summary",
            use_container_width=True,
            key="summary_btn"
        ):

            try:

                with st.spinner("Generating summary..."):

                    result = pipeline.generate_summary(
                        selected_video_id
                    )

                if result.get("success"):

                    st.success("Summary generated.")

                    st.markdown(result["summary"])

                else:
                    show_result_error(result)

            except Exception as error:
                st.error(f"Unexpected error: {error}")


# ==================================================
# TAB 4: STUDY NOTES
# ==================================================

with tab4:

    st.header("Generate Study Notes")

    st.caption("Create clear study notes for revision.")

    videos = get_processed_videos()

    video_options = {
        get_video_id(video): get_video_title(video)
        for video in videos
        if get_video_id(video)
    }

    if not video_options:

        st.warning(
            "No videos processed yet. Go to the Process Video tab first."
        )

    else:

        selected_video_id = st.selectbox(
            "Select a video",
            options=list(video_options.keys()),
            format_func=lambda video_id: video_options[video_id],
            key="notes_video_select"
        )

        if st.button(
            "Generate Study Notes",
            use_container_width=True,
            key="notes_btn"
        ):

            try:

                with st.spinner("Generating study notes..."):

                    result = pipeline.generate_study_notes(
                        selected_video_id
                    )

                if result.get("success"):

                    st.success("Study notes generated.")

                    st.markdown(result["study_notes"])

                    study_notes = result["study_notes"]
                    pdf_data = create_pdf(study_notes)
                    docx_data = create_docx(study_notes)

                    col1, col2, col3 = st.columns(3)
                    with col1:
                        st.download_button(
                            label="Download TXT",
                            data=study_notes,
                            file_name="study_notes.txt",
                            mime="text/plain",
                            key="download_notes_txt"
                        )
                    with col2:
                        st.download_button(
                            label="Download PDF",
                            data=pdf_data,
                            file_name="study_notes.pdf",
                            mime="application/pdf",
                            key="download_notes_pdf"
                        )
                    with col3:
                        st.download_button(
                            label="Download Word",
                            data=docx_data,
                            file_name="study_notes.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key="download_notes_docx"
                        )

                else:
                    show_result_error(result)

            except Exception as error:
                st.error(f"Unexpected error: {error}")


# --------------------------------------------------
# FOOTER
# --------------------------------------------------

st.divider()

st.caption(
    "VideoMind AI | Powered by Streamlit, ChromaDB, "
    "Sentence Transformers, SQLite, and Groq"
)